from __future__ import annotations

import io
import zipfile

from conftest import Client
from docx import Document
from test_files import upload

from studio import export
from studio.markdown import render_html

SAMPLE = """# Plan

Some ==highlighted== text with **bold**, *italic*, ~~gone~~ and `code`.
See [the site](https://example.com).

| Item | Qty |
| :--- | ---: |
| Pi 5 | 3 |

```python
def hello():
    return "hi"
```

1. first
2. second
   - nested

3. third

- [x] done
- [ ] todo

> A quote

A footnote[^1].

---

![logo](/api/files/{image}/content) ![far](https://example.com/x.png)

[^1]: The note.
"""

# A 4x3 PNG.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000040000000308020000003b963991000000144944415478"
    "9c633c99e2c400034c0c4800850300362e01757702fcff0000000049454e44ae426082"
)


def test_markdown_matches_the_editor():
    html = render_html("a ==b== `c` <b>x</b>")
    assert "<mark>b</mark>" in html
    assert "&lt;b&gt;" in html


def note_with_image(client):
    image = upload(client, "logo.png", PNG).json()["file"]
    note = client.post_json(
        "/api/files", {"name": "Plan.md", "content": SAMPLE.replace("{image}", str(image["id"]))}
    )
    return note.json()["file"], image


def test_word_export(user_client):
    note, _ = note_with_image(user_client)
    response = user_client.get(f"/api/files/{note['id']}/export?format=docx")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == "attachment; filename*=UTF-8''Plan.docx"
    doc = Document(io.BytesIO(response.content))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Plan"
    assert doc.paragraphs[0].style.name == "Heading 1"
    body = doc.paragraphs[1]
    highlighted = [r for r in body.runs if r.font.highlight_color is not None]
    assert [r.text for r in highlighted] == ["highlighted"]
    assert any(r.bold and r.text == "bold" for r in body.runs)
    assert any(r.font.name == "Consolas" and r.text == "code" for r in body.runs)
    assert "the site" in body._p.xml and "w:hyperlink" in body._p.xml
    table = doc.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["Pi 5", "3"]
    assert 'def hello():\n    return "hi"' in texts
    assert "☑ done" in texts and "☐ todo" in texts
    numbered = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in numbered] == ["first", "second", "third"]
    assert any(p.style.name == "Quote" for p in doc.paragraphs)
    # The stored image is embedded; the outside one is only named.
    with zipfile.ZipFile(io.BytesIO(response.content)) as z:
        assert any(name.startswith("word/media/") for name in z.namelist())
    assert any("[far]" in t for t in texts)


def test_pdf_export(user_client):
    note, _ = note_with_image(user_client)
    response = user_client.get(f"/api/files/{note['id']}/export?format=pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_pdf_never_fetches_outside_urls(monkeypatch):
    seen = []
    monkeypatch.setattr(export, "internal_file_id", lambda url: seen.append(url) or None)
    pdf = export.to_pdf("![x](http://192.168.1.1/admin.png)", "t", lambda _id: None)
    assert pdf.startswith(b"%PDF")
    assert seen == ["http://192.168.1.1/admin.png"]


def test_export_needs_access_and_markdown(app_factory):
    app = app_factory()
    with Client(app) as kev, Client(app) as other:
        kev.register("kev")
        other.register("other")
        note = kev.post_json("/api/files", {"name": "a.md", "content": "hi"}).json()["file"]
        assert other.get(f"/api/files/{note['id']}/export?format=pdf").status_code == 404
        text = kev.post_json("/api/files", {"name": "a.txt"}).json()["file"]
        assert kev.get(f"/api/files/{text['id']}/export?format=pdf").status_code == 415
        assert kev.get(f"/api/files/{note['id']}/export?format=odt").status_code == 422


def test_images_from_other_users_are_left_out(app_factory):
    app = app_factory()
    with Client(app) as kev, Client(app) as other:
        kev.register("kev")
        other.register("other")
        image = upload(other, "secret.png", PNG).json()["file"]
        note = kev.post_json(
            "/api/files", {"name": "a.md", "content": f"![s](/api/files/{image['id']}/content)"}
        ).json()["file"]
        response = kev.get(f"/api/files/{note['id']}/export?format=docx")
        with zipfile.ZipFile(io.BytesIO(response.content)) as z:
            assert not any(name.startswith("word/media/") for name in z.namelist())
