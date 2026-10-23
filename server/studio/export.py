"""Word (.docx) and PDF export of Markdown notes.

Both start from the same parse as the preview (:mod:`studio.markdown`), so
tables, fenced code with syntax colours, task lists, footnotes and
``==highlights==`` come through.

Images: only pictures stored in Markdown Studio are included (written as
``![alt](/api/files/12/content)``), and only if the reader may see them. The
server never fetches outside URLs while exporting, so a note can't make it
reach into your network; such images become their alt text.
"""

from __future__ import annotations

import html
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from markdown_it.token import Token
from pygments import lex
from pygments.lexers import get_lexer_by_name
from pygments.styles import get_style_by_name
from pygments.util import ClassNotFound

from .markdown import CODE_STYLE, code_css, parser, render_html

# Returns (bytes, mime type) for an image the reader may see, else None.
ImageLoader = Callable[[int], tuple[bytes, str] | None]

CHECKBOX = re.compile(r'<input class="task-list-item-checkbox"[^>]*>')
INTERNAL_IMAGE = re.compile(r"^(?:https?://[^/]+)?/api/files/(\d+)/content(?:\?.*)?$")
MONO = "Consolas"
CODE_BG = "F0EEE6"
MAX_IMAGE_WIDTH = Inches(6)


def internal_file_id(src: str) -> int | None:
    match = INTERNAL_IMAGE.match(src.strip())
    return int(match.group(1)) if match else None


# ------------------------------------------------------------------- PDF

PRINT_CSS = """
@page {
  size: A4;
  margin: 22mm 20mm 24mm;
  @bottom-center { content: counter(page) " / " counter(pages); font: 9pt sans-serif; color: #888; }
}
html { font: 10.5pt/1.55 "DejaVu Sans", "Helvetica", sans-serif; color: #141413; }
h1, h2, h3, h4 { font-family: "DejaVu Serif", Georgia, serif; font-weight: normal;
  line-height: 1.25; margin: 1.4em 0 0.5em; page-break-after: avoid; }
h1 { font-size: 22pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; font-weight: bold; }
body > :first-child { margin-top: 0; }
p, ul, ol, blockquote, pre, table { margin: 0 0 0.8em; }
a { color: #c96442; text-decoration: none; }
mark { background: #fbe7a1; color: inherit; padding: 0 1pt; }
code { font: 9pt "DejaVu Sans Mono", monospace; background: #f0eee6; padding: 0 2pt;
  border-radius: 2pt; }
pre { background: #f0eee6; padding: 8pt 10pt; border-radius: 6pt; white-space: pre-wrap;
  page-break-inside: avoid; }
pre code { background: none; padding: 0; }
blockquote { margin-left: 0; padding-left: 10pt; border-left: 2.5pt solid #e3e0d6; color: #6b6a63; }
table { border-collapse: collapse; }
th, td { border: 0.75pt solid #d6d3c8; padding: 3pt 7pt; }
th { background: #f0eee6; }
img { max-width: 100%; }
hr { border: 0; border-top: 0.75pt solid #d6d3c8; }
.contains-task-list { list-style: none; padding-left: 0.3em; }
.check { margin-right: 4pt; }
.footnotes { font-size: 9pt; color: #6b6a63; }
"""


def to_pdf(source: str, title: str, load_image: ImageLoader) -> bytes:
    # WeasyPrint is slow to import, so only when someone exports a PDF.
    from weasyprint import CSS, HTML
    from weasyprint.urls import URLFetcher, URLFetcherResponse

    class Fetcher(URLFetcher):
        """Stored images the reader may see, and data: URLs. Nothing else."""

        def fetch(self, url, headers=None):
            file_id = internal_file_id(url)
            if file_id is not None:
                image = load_image(file_id)
                if image is not None:
                    return URLFetcherResponse(url, image[0], {"Content-Type": image[1]})
            elif url.startswith("data:"):
                return super().fetch(url, headers)
            raise ValueError("Exports only include images stored in Markdown Studio")

    body = CHECKBOX.sub(
        lambda m: f"<span class='check'>{'☑' if 'checked' in m.group(0) else '☐'}</span>",
        render_html(source),
    )
    document = (
        f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
        f"</head><body>{body}</body></html>"
    )
    return HTML(
        string=document, base_url="https://studio.invalid/", url_fetcher=Fetcher()
    ).write_pdf(stylesheets=[CSS(string=PRINT_CSS + code_css())])


# ------------------------------------------------------------------ Word


@dataclass
class _Lists:
    """Open lists, innermost last: [ordered, start, numbering id or None, is a task list]."""

    stack: list[list] = field(default_factory=list)
    item_started: list[bool] = field(default_factory=list)


class _DocxWriter:
    def __init__(self, title: str, load_image: ImageLoader) -> None:
        self.doc = Document()
        self.doc.core_properties.title = title
        self.load_image = load_image
        self.lists = _Lists()
        self.quote_depth = 0
        self.paragraph = None
        self.footnote_label: str | None = None
        self.code_colors = _code_colors()
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        for level, size in enumerate((22, 16, 13, 12, 11, 11), start=1):
            heading = self.doc.styles[f"Heading {level}"]
            heading.font.name = "Georgia"
            heading.font.size = Pt(size)
            heading.font.bold = level >= 3
            heading.font.color.rgb = RGBColor(0x14, 0x14, 0x13)

    # -- blocks

    def write(self, tokens: list[Token]) -> bytes:
        i = 0
        while i < len(tokens):
            i = self.block(tokens, i)
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()

    def block(self, tokens: list[Token], i: int) -> int:
        token = tokens[i]
        kind = token.type
        if kind == "heading_open":
            level = int(token.tag[1])
            self.paragraph = self.doc.add_heading(level=min(level, 9))
            self.inline(tokens[i + 1])
            return i + 3
        if kind == "paragraph_open":
            self.paragraph = self.new_paragraph()
            if self.footnote_label:
                run = self.paragraph.add_run(f"{self.footnote_label} ")
                run.font.superscript = True
                self.footnote_label = None
            self.inline(tokens[i + 1])
            return i + 3
        if kind in ("bullet_list_open", "ordered_list_open"):
            ordered = kind == "ordered_list_open"
            start = int(token.attrs.get("start", 1)) if ordered else 1
            tasks = "contains-task-list" in str(token.attrs.get("class", ""))
            self.lists.stack.append([ordered, start, None, tasks])
            return i + 1
        if kind in ("bullet_list_close", "ordered_list_close"):
            self.lists.stack.pop()
            return i + 1
        if kind == "list_item_open":
            self.lists.item_started.append(False)
            return i + 1
        if kind == "list_item_close":
            self.lists.item_started.pop()
            return i + 1
        if kind == "blockquote_open":
            self.quote_depth += 1
            return i + 1
        if kind == "blockquote_close":
            self.quote_depth -= 1
            return i + 1
        if kind in ("fence", "code_block"):
            self.code_block(token.content, token.info.strip().split(" ")[0] if token.info else "")
            return i + 1
        if kind == "hr":
            self.rule()
            return i + 1
        if kind == "table_open":
            return self.table(tokens, i)
        if kind == "footnote_block_open":
            self.rule()
            return i + 1
        if kind == "footnote_open":
            self.footnote_label = str(token.meta.get("id", 0) + 1)
            return i + 1
        return i + 1  # closers and anything without its own output

    def new_paragraph(self):
        if self.lists.stack:
            depth = len(self.lists.stack)
            ordered, start, num_id, tasks = self.lists.stack[-1]
            if not self.lists.item_started[-1] and not tasks:
                self.lists.item_started[-1] = True
                base = "List Number" if ordered else "List Bullet"
                style = base if depth == 1 else f"{base} {min(depth, 3)}"
                paragraph = self.doc.add_paragraph(style=style)
                if ordered:
                    if num_id is None:
                        num_id = _restart_numbering(self.doc, paragraph, start)
                        self.lists.stack[-1][2] = num_id
                    else:
                        _set_num_id(paragraph, num_id)
                return paragraph
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25 + 0.25 * depth)
            return paragraph
        if self.quote_depth:
            paragraph = self.doc.add_paragraph(style="Quote")
            paragraph.paragraph_format.left_indent = Inches(0.4 * self.quote_depth)
            return paragraph
        return self.doc.add_paragraph()

    def code_block(self, code: str, lang: str) -> None:
        paragraph = self.doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(4)
        fmt.space_after = Pt(8)
        if self.lists.stack:
            fmt.left_indent = Inches(0.25 + 0.25 * len(self.lists.stack))
        _shade(paragraph._p.get_or_add_pPr(), CODE_BG)
        code = code.rstrip("\n")
        try:
            lexer = get_lexer_by_name(lang, ensurenl=False, stripnl=False) if lang else None
        except ClassNotFound:
            lexer = None
        pieces = lex(code, lexer) if lexer else [(None, code)]
        for token_type, value in pieces:
            lines = value.split("\n")
            for n, line in enumerate(lines):
                if n:
                    paragraph.add_run().add_break()
                if line:
                    run = paragraph.add_run(line)
                    run.font.name = MONO
                    run.font.size = Pt(9.5)
                    color = self.code_colors(token_type)
                    if color:
                        run.font.color.rgb = RGBColor.from_string(color)

    def rule(self) -> None:
        paragraph = self.doc.add_paragraph()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "D6D3C8")):
            bottom.set(qn(f"w:{key}"), value)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)

    def table(self, tokens: list[Token], i: int) -> int:
        rows: list[list[tuple[Token, str]]] = []
        while tokens[i].type != "table_close":
            token = tokens[i]
            if token.type == "tr_open":
                rows.append([])
            elif token.type in ("th_open", "td_open"):
                style = str(token.attrs.get("style", ""))
                align = style.removeprefix("text-align:") if style.startswith("text-align:") else ""
                rows[-1].append((tokens[i + 1], align))
            i += 1
        columns = max((len(r) for r in rows), default=0)
        if not columns:
            return i + 1
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for r, row in enumerate(rows):
            for c, (inline, align) in enumerate(row):
                cell = table.cell(r, c)
                self.paragraph = cell.paragraphs[0]
                self.paragraph.alignment = {
                    "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
                self.inline(inline, bold=r == 0)
                if r == 0:
                    _shade(cell._tc.get_or_add_tcPr(), CODE_BG)
        self.doc.add_paragraph()
        return i + 1

    # -- inline

    def inline(self, token: Token, bold: bool = False) -> None:
        state = {"strong": bold, "em": False, "s": False, "mark": False, "link": None}
        for child in token.children or []:
            kind = child.type
            if kind.endswith("_open") and kind[:-5] in ("strong", "em", "s", "mark"):
                state[kind[:-5]] = True
            elif kind.endswith("_close") and kind[:-6] in ("strong", "em", "s", "mark"):
                state[kind[:-6]] = kind[:-6] == "strong" and bold
            elif kind == "link_open":
                state["link"] = str(child.attrs.get("href", ""))
            elif kind == "link_close":
                state["link"] = None
            elif kind == "text":
                self.text(child.content, state)
            elif kind == "code_inline":
                run = self.text(child.content, state)
                run.font.name = MONO
                run.font.size = Pt(9.5)
                _shade(run._r.get_or_add_rPr(), CODE_BG)
            elif kind == "softbreak":
                self.text(" ", state)
            elif kind == "hardbreak":
                self.paragraph.add_run().add_break()
            elif kind == "html_inline" and "task-list-item-checkbox" in child.content:
                self.text("☑" if "checked" in child.content else "☐", state)
            elif kind == "html_inline":
                self.text(child.content, state)  # raw HTML is shown, not run
            elif kind == "image":
                self.image(child, state)
            elif kind == "footnote_ref":
                run = self.paragraph.add_run(str(child.meta.get("id", 0) + 1))
                run.font.superscript = True
            elif kind == "footnote_anchor":
                continue

    def text(self, value: str, state: dict):
        if state["link"]:
            return _add_hyperlink(self.paragraph, state["link"], value, state)
        run = self.paragraph.add_run(value)
        _format(run, state)
        return run

    def image(self, token: Token, state: dict) -> None:
        alt = token.content or "".join(c.content for c in token.children or [])
        file_id = internal_file_id(str(token.attrs.get("src", "")))
        loaded = self.load_image(file_id) if file_id is not None else None
        if loaded is not None:
            try:
                info = DocxImage.from_blob(loaded[0])
                width = Emu(min(int(info.width), MAX_IMAGE_WIDTH))
                self.paragraph.add_run().add_picture(io.BytesIO(loaded[0]), width=width)
                return
            except Exception:  # noqa: S110 (not an image Word understands: fall back to alt text)
                pass
        run = self.paragraph.add_run(f"[{alt or 'image'}]")
        run.italic = True


def _format(run, state: dict) -> None:
    run.bold = state["strong"] or None
    run.italic = state["em"] or None
    if state["s"]:
        run.font.strike = True
    if state["mark"]:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _add_hyperlink(paragraph, url: str, text: str, state: dict):
    if url.startswith("#"):
        run = paragraph.add_run(text)
        _format(run, state)
        return run
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = paragraph.add_run(text)
    _format(run, state)
    run.font.color.rgb = RGBColor(0xC9, 0x64, 0x42)
    run.font.underline = True
    run._r.getparent().remove(run._r)
    link.append(run._r)
    paragraph._p.append(link)
    return run


def _shade(properties, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_num_id(paragraph, num_id: int) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def _restart_numbering(doc, paragraph, start: int) -> int:
    """Give this ordered list its own numbering, so it starts at ``start``."""
    numbering = doc.part.numbering_part.element
    style = doc.styles[paragraph.style.name]
    style_num = style.element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(start)
    _set_num_id(paragraph, num.numId)
    return num.numId


def _code_colors():
    style = get_style_by_name(CODE_STYLE)
    cache: dict = {}

    def color(token_type) -> str | None:
        if token_type is None:
            return None
        if token_type not in cache:
            cache[token_type] = style.style_for_token(token_type).get("color") or None
        return cache[token_type]

    return color


def to_docx(source: str, title: str, load_image: ImageLoader) -> bytes:
    return _DocxWriter(title, load_image).write(parser().parse(source))
